"""
Text extraction service for PDF, DOCX, and TXT files.
"""

import io
import logging
from typing import Optional

from docx import Document as DocxDocument
from PyPDF2 import PdfReader

logger = logging.getLogger(__name__)


class TextExtractionError(Exception):
    """Raised when text extraction fails."""

    pass


def extract_text_from_pdf(file_content: bytes) -> tuple[str, Optional[int]]:
    """
    Extract text from PDF file.

    Args:
        file_content: Raw PDF file content

    Returns:
        Tuple of (extracted_text, page_count)

    Raises:
        TextExtractionError: If extraction fails
    """
    try:
        pdf_reader = PdfReader(io.BytesIO(file_content))
        page_count = len(pdf_reader.pages)

        text_parts = []
        for page_num, page in enumerate(pdf_reader.pages, 1):
            try:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(f"--- Page {page_num} ---\n{page_text}")
            except Exception as e:
                logger.warning(f"Failed to extract text from PDF page {page_num}: {e}")
                # Continue with next page

        extracted_text = "\n\n".join(text_parts)

        if not extracted_text.strip():
            raise TextExtractionError("PDF contains no extractable text")

        return extracted_text, page_count

    except Exception as e:
        raise TextExtractionError(f"Failed to extract text from PDF: {str(e)}")


def extract_text_from_docx(file_content: bytes) -> tuple[str, None]:
    """
    Extract text from DOCX file.

    Args:
        file_content: Raw DOCX file content

    Returns:
        Tuple of (extracted_text, None) — DOCX doesn't have pages

    Raises:
        TextExtractionError: If extraction fails
    """
    try:
        docx = DocxDocument(io.BytesIO(file_content))

        # Extract text from paragraphs
        paragraphs = []
        for para in docx.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)

        # Extract text from tables
        for table in docx.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                if any(row_cells):
                    paragraphs.append(" | ".join(row_cells))

        extracted_text = "\n\n".join(paragraphs)

        if not extracted_text.strip():
            raise TextExtractionError("DOCX contains no extractable text")

        return extracted_text, None

    except Exception as e:
        raise TextExtractionError(f"Failed to extract text from DOCX: {str(e)}")


def extract_text_from_txt(file_content: bytes) -> tuple[str, None]:
    """
    Extract text from TXT file.

    Args:
        file_content: Raw TXT file content

    Returns:
        Tuple of (extracted_text, None)

    Raises:
        TextExtractionError: If extraction fails
    """
    try:
        # Try UTF-8 first, then fallback to latin-1
        try:
            text = file_content.decode("utf-8")
        except UnicodeDecodeError:
            text = file_content.decode("latin-1")

        if not text.strip():
            raise TextExtractionError("TXT file is empty")

        return text, None

    except Exception as e:
        raise TextExtractionError(f"Failed to extract text from TXT: {str(e)}")


def extract_text(
    file_content: bytes, file_type: str
) -> tuple[str, Optional[int]]:
    """
    Extract text from file based on type.

    Args:
        file_content: Raw file content
        file_type: File type ('pdf', 'docx', 'txt')

    Returns:
        Tuple of (extracted_text, page_count or None)

    Raises:
        TextExtractionError: If extraction fails
    """
    if file_type == "pdf":
        return extract_text_from_pdf(file_content)
    elif file_type == "docx":
        return extract_text_from_docx(file_content)
    elif file_type == "txt":
        return extract_text_from_txt(file_content)
    else:
        raise TextExtractionError(f"Unsupported file type: {file_type}")
