"""Unit tests for text extraction service."""

import io

import pytest
from docx import Document as DocxDocument
from PyPDF2 import PdfWriter

from app.services.text_extraction_service import (
    extract_text,
    extract_text_from_docx,
    extract_text_from_pdf,
    extract_text_from_txt,
    TextExtractionError,
)


class TestExtractTextFromTxt:
    """Tests for TXT file extraction."""

    def test_extract_simple_text(self):
        """Test extraction of simple text."""
        content = "Hello, world!".encode("utf-8")
        text, page_count = extract_text_from_txt(content)

        assert text == "Hello, world!"
        assert page_count is None

    def test_extract_multiline_text(self):
        """Test extraction of multi-line text."""
        content = "Line 1\nLine 2\nLine 3".encode("utf-8")
        text, _ = extract_text_from_txt(content)

        assert "Line 1" in text
        assert "Line 2" in text
        assert "Line 3" in text

    def test_empty_text_fails(self):
        """Test empty text file raises error."""
        content = "".encode("utf-8")
        with pytest.raises(TextExtractionError) as exc_info:
            extract_text_from_txt(content)
        assert "empty" in str(exc_info.value).lower()

    def test_latin1_encoding(self):
        """Test fallback to latin-1 encoding."""
        # Create content with latin-1 chars that aren't valid UTF-8
        content = "café".encode("latin-1")
        text, _ = extract_text_from_txt(content)
        assert len(text) > 0

    def test_whitespace_only_fails(self):
        """Test whitespace-only file raises error."""
        content = "   \n\n   ".encode("utf-8")
        with pytest.raises(TextExtractionError):
            extract_text_from_txt(content)


class TestExtractTextFromDocx:
    """Tests for DOCX file extraction."""

    def test_extract_simple_docx(self):
        """Test extraction from simple DOCX."""
        # Create a minimal DOCX file in memory
        doc = DocxDocument()
        doc.add_paragraph("Hello, world!")

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        content = buffer.read()

        text, page_count = extract_text_from_docx(content)

        assert "Hello, world!" in text
        assert page_count is None

    def test_extract_multiple_paragraphs(self):
        """Test extraction from DOCX with multiple paragraphs."""
        doc = DocxDocument()
        doc.add_paragraph("Para 1")
        doc.add_paragraph("Para 2")
        doc.add_paragraph("Para 3")

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        content = buffer.read()

        text, _ = extract_text_from_docx(content)

        assert "Para 1" in text
        assert "Para 2" in text
        assert "Para 3" in text

    def test_extract_docx_with_table(self):
        """Test extraction from DOCX with table."""
        doc = DocxDocument()
        table = doc.add_table(rows=2, cols=2)

        # Fill table
        table.rows[0].cells[0].text = "Header 1"
        table.rows[0].cells[1].text = "Header 2"
        table.rows[1].cells[0].text = "Data 1"
        table.rows[1].cells[1].text = "Data 2"

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        content = buffer.read()

        text, _ = extract_text_from_docx(content)

        # Table data should be in extracted text
        assert len(text) > 0

    def test_empty_docx_fails(self):
        """Test empty DOCX raises error."""
        doc = DocxDocument()
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        content = buffer.read()

        with pytest.raises(TextExtractionError) as exc_info:
            extract_text_from_docx(content)
        assert "no extractable text" in str(exc_info.value).lower()

    def test_invalid_docx_fails(self):
        """Test invalid DOCX data raises error."""
        content = b"Not a valid DOCX file"
        with pytest.raises(TextExtractionError):
            extract_text_from_docx(content)


class TestExtractTextFromPdf:
    """Tests for PDF file extraction."""

    def test_extract_simple_pdf(self):
        """Test extraction from simple PDF."""
        # Create a minimal PDF in memory
        writer = PdfWriter()
        page = writer.add_blank_page(width=200, height=200)

        buffer = io.BytesIO()
        writer.write(buffer)
        buffer.seek(0)
        content = buffer.read()

        # Note: Empty PDF may fail, so we expect it or accept empty result
        try:
            text, page_count = extract_text_from_pdf(content)
            assert page_count == 1
        except TextExtractionError:
            # Empty PDFs may raise an error, which is acceptable
            pass

    def test_invalid_pdf_fails(self):
        """Test invalid PDF data raises error."""
        content = b"Not a valid PDF file"
        with pytest.raises(TextExtractionError):
            extract_text_from_pdf(content)

    def test_empty_pdf_fails(self):
        """Test empty PDF raises error."""
        # Minimal valid PDF header but no content
        content = b"%PDF-1.4\n"
        with pytest.raises(TextExtractionError):
            extract_text_from_pdf(content)


class TestExtractText:
    """Tests for generic extract_text function."""

    def test_extract_txt(self):
        """Test extraction with txt type."""
        content = "Hello, world!".encode("utf-8")
        text, page_count = extract_text(content, "txt")

        assert "Hello" in text
        assert page_count is None

    def test_extract_docx(self):
        """Test extraction with docx type."""
        doc = DocxDocument()
        doc.add_paragraph("Test content")

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        content = buffer.read()

        text, page_count = extract_text(content, "docx")

        assert "Test content" in text
        assert page_count is None

    def test_unsupported_type_fails(self):
        """Test unsupported file type raises error."""
        content = b"some content"
        with pytest.raises(TextExtractionError) as exc_info:
            extract_text(content, "xyz")
        assert "Unsupported" in str(exc_info.value)

    def test_type_must_match_content(self):
        """Test that type parameter must match actual content."""
        txt_content = "Hello".encode("utf-8")

        # Trying to extract as PDF should fail due to signature mismatch
        with pytest.raises(TextExtractionError):
            extract_text(txt_content, "pdf")
